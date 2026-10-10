"""Consolidated test suite for document parsers and extraction pipelines.

Covers:
- Document format parsing (PDF with OCR fallback, DOCX with tables, EML emails)
- Deterministic regex & NER extraction (emails, phones, skills, offsets)
- Hybrid candidate profile extraction & validation (confidence scoring, name/email matching, AI fallback)
"""
import io
from pathlib import Path
from typing import Any, Dict, List
import pytest

from ingestion.parsers.docx_parser import parse_docx
from ingestion.parsers.email_parser import parse_email
from ingestion.parsers.models import ParsedDocument
from ingestion.parsers.pdf_parser import parse_pdf
from candidate_intelligence_platform.extraction.deterministic_ner import extract_facts, get_nlp
from candidate_intelligence_platform.extraction.hybrid_extractor import (
    calculate_tier1_confidence,
    extract_candidate_profile_hybrid,
    validate_name_against_email,
)


# ---------------------------------------------------------------------------
# Document Generation Helpers
# ---------------------------------------------------------------------------

def _make_minimal_pdf(text: str = "John Doe\nSoftware Engineer") -> bytes:
    """Build a minimal valid single-page PDF containing text."""
    safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content_stream = f"BT\n/F1 12 Tf\n50 700 Td\n({safe_text}) Tj\nET".encode()
    stream_len = len(content_stream)

    pdf = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R "
        b"/MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
        + f"4 0 obj\n<< /Length {stream_len} >>\nstream\n".encode()
        + content_stream
        + b"\nendstream\nendobj\n"
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    )

    xref_offset = len(pdf)
    offsets = []
    pos = 0
    for line in pdf.split(b"\n"):
        if line.endswith(b"obj"):
            offsets.append(pos)
        pos += len(line) + 1

    xref = b"xref\n" + f"0 {len(offsets) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        xref += f"{off:010d} 00000 n \n".encode()

    trailer = f"trailer\n<< /Size {len(offsets) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return pdf + xref + trailer


def _make_docx(tmp_path: Path, paragraphs: List[str], table_data: List[List[str]] | None = None) -> Path:
    from docx import Document

    doc = Document()
    for para in paragraphs:
        doc.add_paragraph(para)

    if table_data:
        rows = len(table_data)
        cols = len(table_data[0]) if rows else 0
        table = doc.add_table(rows=rows, cols=cols)
        for r_idx, row_cells in enumerate(table_data):
            for c_idx, cell_text in enumerate(row_cells):
                table.rows[r_idx].cells[c_idx].text = cell_text

    docx_path = tmp_path / "resume.docx"
    doc.save(str(docx_path))
    return docx_path


def _make_eml(tmp_path: Path, subject: str = "Resume", body: str = "Please find resume attached.") -> Path:
    eml_path = tmp_path / "message.eml"
    content = (
        f"From: candidate@example.com\r\n"
        f"To: recruiter@company.com\r\n"
        f"Subject: {subject}\r\n"
        f"MIME-Version: 1.0\r\n"
        f"Content-Type: text/plain; charset=utf-8\r\n"
        f"\r\n"
        f"{body}\r\n"
    )
    eml_path.write_bytes(content.encode("utf-8"))
    return eml_path


# ---------------------------------------------------------------------------
# Format Parsers: PDF, DOCX, EML
# ---------------------------------------------------------------------------

def test_parse_pdf_functional_flow(tmp_path: Path) -> None:
    content = "Jane Smith Senior Developer"
    pdf_file = tmp_path / "resume.pdf"
    pdf_file.write_bytes(_make_minimal_pdf(content))

    result = parse_pdf(pdf_file)
    assert isinstance(result, ParsedDocument)
    assert "Jane Smith" in result.text
    assert result.pages == 1
    assert "source_path" in result.metadata


def test_parse_pdf_ocr_fallback(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import fitz
    from ingestion.parsers import pdf_parser

    monkeypatch.setattr(fitz.Page, "get_text", lambda self, *args, **kwargs: "")

    # Mock pdfplumber to return empty text
    monkeypatch.setattr(pdf_parser, "_extract_with_pdfplumber", lambda p, n: "")

    try:
        import pytesseract
        monkeypatch.setattr(pytesseract, "image_to_string", lambda img: "Mocked OCR Text")
    except ImportError:
        pytest.skip("pytesseract not installed")

    pdf_file = tmp_path / "scanned.pdf"
    pdf_file.write_bytes(_make_minimal_pdf("Hidden"))

    result = parse_pdf(pdf_file)
    assert isinstance(result, ParsedDocument)
    assert "Mocked OCR Text" in result.text


def test_parse_docx_functional_flow(tmp_path: Path) -> None:
    docx_path = _make_docx(
        tmp_path,
        paragraphs=["Alice Johnson", "Product Manager at Acme Corp", "Skills"],
        table_data=[["Python", "5 years"], ["SQL", "3 years"]],
    )
    result = parse_docx(docx_path)
    assert isinstance(result, ParsedDocument)
    assert "Alice Johnson" in result.text
    assert "Product Manager" in result.text
    assert "Python" in result.text
    assert "SQL" in result.text


def test_parse_email_functional_flow(tmp_path: Path) -> None:
    body = "Please find my resume attached. I have 7 years of Python experience."
    eml_path = _make_eml(tmp_path, subject="Application for SWE Role", body=body)

    result = parse_email(eml_path)
    assert isinstance(result, ParsedDocument)
    assert "resume attached" in result.text
    assert "Python experience" in result.text
    assert result.pages == 1
    assert "Application for SWE Role" in result.metadata["subject"]
    assert "candidate@example.com" in result.metadata["from"]


# ---------------------------------------------------------------------------
# Deterministic NER & Fact Extraction
# ---------------------------------------------------------------------------

def test_extract_email() -> None:
    text = "Contact me at jane.doe@example.com for more info."
    facts = extract_facts(text)
    emails = [f for f in facts if f["claim_category"] == "CONTACT" and f["claim_key"] == "email"]
    assert len(emails) == 1
    assert emails[0]["claim_value"] == "jane.doe@example.com"


def test_extract_phone() -> None:
    text = "Call me: +1 (555) 123-4567"
    facts = extract_facts(text)
    phones = [f for f in facts if f["claim_category"] == "CONTACT" and f["claim_key"] == "phone"]
    assert len(phones) == 1
    assert "555" in phones[0]["claim_value"]


@pytest.mark.skipif(get_nlp() is None, reason="spaCy model en_core_web_sm is not installed")
def test_extract_name() -> None:
    text = "John Doe is a Senior Software Engineer."
    facts = extract_facts(text)
    names = [f for f in facts if f["claim_category"] == "PERSON" and f["claim_key"] == "name"]
    assert len(names) >= 1
    assert names[0]["claim_value"] == "John Doe"


@pytest.mark.skipif(get_nlp() is None, reason="spaCy model en_core_web_sm is not installed")
def test_extract_location() -> None:
    text = "I am based in San Francisco, CA."
    facts = extract_facts(text)
    locations = [f for f in facts if f["claim_category"] == "LOCATION" and f["claim_key"] == "city"]
    assert len(locations) >= 1
    assert "San Francisco" in locations[0]["claim_value"]


def test_extract_skills() -> None:
    text = "I have 5 years of experience in Python and C++."
    facts = extract_facts(text)
    skills = [f for f in facts if f["claim_category"] == "SKILL"]
    skill_values = [s["claim_value"].lower() for s in skills]
    assert "python" in skill_values
    assert "c++" in skill_values


def test_offsets() -> None:
    text = "Email: test@test.com"
    facts = extract_facts(text)
    assert len(facts) > 0
    fact = facts[0]
    start = fact["source_char_offset_start"]
    end = fact["source_char_offset_end"]
    assert text[start:end] == "test@test.com"


# ---------------------------------------------------------------------------
# Hybrid Extraction & Validation
# ---------------------------------------------------------------------------

def test_calculate_tier1_confidence_high() -> None:
    facts = [
        {"claim_category": "PERSON", "claim_key": "name", "claim_value": "John Doe"},
        {"claim_category": "SKILL", "claim_key": "python", "claim_value": "Python"},
    ]
    parsed_profile = {
        "first_name": "John",
        "last_name": "Doe",
        "primary_email": "john.doe@example.com",
        "primary_phone": "555-123-4567",
        "current_title": "Senior Software Engineer",
    }
    score = calculate_tier1_confidence(facts, parsed_profile)
    assert score == 1.0


def test_calculate_tier1_confidence_no_email_phone_bonus() -> None:
    facts = [
        {"claim_category": "CONTACT", "claim_key": "email", "claim_value": "john.doe@example.com"},
        {"claim_category": "CONTACT", "claim_key": "phone", "claim_value": "555-123-4567"},
    ]
    parsed_profile = {
        "first_name": "John",
        "last_name": "Doe",
        "primary_email": "john.doe@example.com",
        "primary_phone": "555-123-4567",
        "current_title": "Candidate",
    }
    score = calculate_tier1_confidence(facts, parsed_profile)
    assert score == 0.50


def test_calculate_tier1_confidence_very_low() -> None:
    facts: List[Dict[str, Any]] = []
    parsed_profile = {
        "first_name": "Uploaded",
        "last_name": "Candidate",
        "primary_email": None,
        "primary_phone": None,
        "current_title": "Candidate",
    }
    score = calculate_tier1_confidence(facts, parsed_profile)
    assert score == 0.0


def test_extract_candidate_profile_hybrid_deterministic() -> None:
    text = """
    John Doe
    john.doe@example.com | (555) 123-4567
    Senior Software Engineer
    San Francisco, CA
    Skills: Python, SQL, Docker
    """
    res = extract_candidate_profile_hybrid(text, confidence_threshold=0.70)
    assert res["first_name"] == "John"
    assert res["last_name"] == "Doe"
    assert res["primary_email"] == "john.doe@example.com"
    assert res["used_ai_fallback"] is False
    assert res["confidence_score"] >= 0.70


def test_extract_candidate_profile_hybrid_ai_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    text = "John Doe\nSenior Software Engineer"

    class MockResponse:
        class message:
            content = '{"claims": [{"claim_category": "PERSON", "claim_key": "name", "claim_value": "John Doe"}, {"claim_category": "CONTACT", "claim_key": "email", "claim_value": "john.doe@inferred.com"}]}'

    monkeypatch.setattr("ollama.chat", lambda *args, **kwargs: MockResponse())
    res = extract_candidate_profile_hybrid(text)
    assert res["used_ai_fallback"] is True
    assert res["primary_email"] == "john.doe@inferred.com"


def test_extract_candidate_profile_hybrid_ai_fallback_failed_warning(monkeypatch: pytest.MonkeyPatch) -> None:
    text = "Jane Smith\nSoftware Engineer"
    monkeypatch.setattr("candidate_intelligence_platform.extraction.hybrid_extractor.extract_inferences", lambda *args, **kwargs: [])
    res = extract_candidate_profile_hybrid(text)
    assert res["used_ai_fallback"] is False
    assert hasattr(res, "warnings") and len(res["warnings"]) > 0
    assert any("LLM fallback attempted" in w for w in res["warnings"])


def test_extract_candidate_profile_with_page_header_noise() -> None:
    text = """
    Page 1 of 2
    SRINIVAS PEDDI
    Alpharetta, GA | (404) 992-1973 | srinipe28@gmail.com
    ____________________________________________________
    PROFESSIONAL SUMMARY
    Cybersecurity professional with extensive experience partnering with engineering teams.
    Skills: Python, SQL, Docker, AWS
    """
    res = extract_candidate_profile_hybrid(text, confidence_threshold=0.70)
    assert res["first_name"] == "Srinivas"
    assert res["last_name"] == "Peddi"
    assert res["primary_email"] == "srinipe28@gmail.com"
    assert res["primary_phone"] == "(404) 992-1973"
    assert "Cybersecurity" in res["current_title"]
    assert res["used_ai_fallback"] is False
    assert res["confidence_score"] >= 0.70


def test_validate_name_against_email() -> None:
    assert validate_name_against_email("John Doe", "john.doe@example.com") is True
    assert validate_name_against_email("John Doe", "johndoe@example.com") is True
    assert validate_name_against_email("John Smith", "jsmith@example.com") is True
    assert validate_name_against_email("John Smith", "j.smith@example.com") is True
    assert validate_name_against_email("Jonathan Doe", "jon.doe@example.com") is True
    assert validate_name_against_email("Jane Smith", None) is True
    assert validate_name_against_email("Jane Smith", "") is True
    assert validate_name_against_email("Executive Summary", "john.doe@example.com") is False
    assert validate_name_against_email("Profile Overview", "jsmith@example.com") is False
    assert validate_name_against_email("John", "alex.smith@example.com") is False
